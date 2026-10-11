from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL = ROOT / "Reglas_de_validacion_OSE_normalizadas.xlsx"
DEFAULT_SOURCE_DIR = ROOT / "rules" / "source"
DEFAULT_GENERATED_DIR = ROOT / "rules" / "generated"
DEFAULT_REPORTS_DIR = ROOT / "reports"
IMPL_FILE = ROOT / "rules_invoice_impl.drl"

TIPOS_DOCUMENTO = [
    "Factura",
    "Boleta",
    "Nota de Crédito",
    "Nota de Débito",
    "Resumen Diario",
    "Bajas",
    "Retenciones",
    "Percepciones",
]


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value or "")
    value = value.lower()
    value = value.replace("_", " ")
    value = value.replace("/", " ")
    value = value.replace("-", " ")
    value = re.sub(r"[^a-z0-9\s]", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def ensure_dirs() -> None:
    DEFAULT_SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    DEFAULT_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    DEFAULT_REPORTS_DIR.mkdir(parents=True, exist_ok=True)


def clean_description(desc: str) -> str:
    desc = " ".join(str(desc).split())
    return desc.replace('"', '\\"')


def render_source_template(codigo: str, desc: str) -> str:
    """Plantilla vacia que genera_source_rules escribe para cada codigo.

    Es la unica definicion de la forma del stub: el reporte la reusa para saber
    si un archivo source sigue sin implementar, en vez de adivinarlo con un regex.
    """
    return f'''rule "R{codigo}"
    when
        $ctx : CpeContext(false)
        $res : ResultadoValidacion()
    then
        $res.addError("{codigo}", "{clean_description(desc)}");
end
'''


def normalize_rule_block(block: str) -> str:
    """Reescribe el nombre de la regla a formato R{code} si empieza con un código numérico."""
    name_match = re.match(r'^rule\s+"(\d{4})_(.*)"', block, flags=re.M)
    if name_match:
        code = name_match.group(1)
        return re.sub(
            r'^rule\s+"\d{4}_.*"', f'rule "R{code}"', block, count=1, flags=re.M
        )
    return block


def load_source_stubs(excel_path: Path = DEFAULT_EXCEL) -> Dict[str, str]:
    df = pd.read_excel(excel_path, sheet_name="CódigosRetorno")
    stubs: Dict[str, str] = {}
    for _, row in df.iterrows():
        codigo = str(row["Código"]).strip().zfill(4)
        desc = str(row["Descripción de Error u Observación"]).strip()
        stubs[codigo] = render_source_template(codigo, desc)
    return stubs


def generate_source_rules(excel_path: Path = DEFAULT_EXCEL, source_dir: Path = DEFAULT_SOURCE_DIR) -> Dict[str, int]:
    ensure_dirs()
    df = pd.read_excel(excel_path, sheet_name="CódigosRetorno")
    created_count = 0
    updated_count = 0
    for _, row in df.iterrows():
        codigo = str(row["Código"]).strip().zfill(4)
        desc = str(row["Descripción de Error u Observación"]).strip()
        template = render_source_template(codigo, desc)
        file_path = source_dir / f"R{codigo}.drl"
        if not file_path.exists():
            file_path.write_text(template, encoding="utf-8")
            created_count += 1
        else:
            updated_count += 1
    return {"created": created_count, "updated": updated_count, "total": len(list(source_dir.glob("R*.drl")))}


def generate_generated_files(source_dir: Path = DEFAULT_SOURCE_DIR, generated_dir: Path = DEFAULT_GENERATED_DIR) -> Dict[str, int]:
    generated_dir.mkdir(parents=True, exist_ok=True)
    buffers: Dict[str, List[str]] = {tipo: [] for tipo in TIPOS_DOCUMENTO}
    df = pd.read_excel(DEFAULT_EXCEL, sheet_name="CódigosRetorno")
    for _, row in df.iterrows():
        codigo = str(row["Código"]).zfill(4)
        regla_path = source_dir / f"R{codigo}.drl"
        if not regla_path.exists():
            continue
        regla_texto = regla_path.read_text(encoding="utf-8")
        tipos_regla = [t.strip() for t in str(row["Tipo(s) de documento"]).split(";") if t.strip()]
        for tipo in tipos_regla:
            if tipo in buffers:
                buffers[tipo].append(f"// Código {codigo}\n{regla_texto}\n")
    written: Dict[str, int] = {}
    for tipo, reglas in buffers.items():
        normalized_name = tipo.lower().replace(" ", "_").replace("é", "e").replace("ó", "o").replace("á", "a").replace("í", "i").replace("ú", "u")
        output = generated_dir / f"{normalized_name}.drl"
        output.write_text("\n".join(reglas), encoding="utf-8")
        written[tipo] = len(reglas)
    return written


def extract_rule_name(block: str) -> str:
    match = re.search(r'^rule\s+"([^"]+)"', block, flags=re.M)
    return match.group(1) if match else ""


def normalize_rule_name(name: str) -> str:
    """Si el nombre empieza con un código numérico como '1001_...', lo transforma a 'R1001'."""
    match = re.match(r'^(\d{4})_(.*)', name)
    if match:
        return f"R{match.group(1)}"
    return name


def load_code_to_doc_types(excel_path: Path = DEFAULT_EXCEL) -> Dict[str, set]:
    """Carga el mapeo de códigos a tipos de documento desde el Excel."""
    df = pd.read_excel(excel_path, sheet_name="CódigosRetorno")
    code_to_docs: Dict[str, set] = defaultdict(set)
    for _, row in df.iterrows():
        codigo = str(row["Código"]).strip().zfill(4)
        tipos_str = str(row.get("Tipo(s) de documento", ""))
        for tipo in tipos_str.split(";"):
            tipo = tipo.strip()
            if tipo:
                code_to_docs[codigo].add(tipo)
    return code_to_docs


def extract_rule_code(block: str) -> Optional[str]:
    match = re.search(r'add(?:Error|Observacion|Excepcion)\("(\d{4})"', block)
    return match.group(1) if match else None


def extract_description(block: str) -> str:
    match = re.search(r'add(?:Error|Observacion|Excepcion)\("\d{4}"\s*,\s*"([^"]*)"\)', block)
    return match.group(1) if match else ""


def extract_when_condition(block: str) -> str:
    match = re.search(r'(?is)\bwhen\b(.*?)(?:\bthen\b|$)', block)
    if not match:
        return ""
    return normalize_text(match.group(1))


def parse_rule_block(block: str) -> Dict[str, str]:
    raw_name = extract_rule_name(block)
    return {
        "name": normalize_rule_name(raw_name),
        "original_name": raw_name,
        "code": extract_rule_code(block) or "",
        "description": normalize_text(extract_description(block)),
        "condition": extract_when_condition(block),
    }


def parse_impl_rules(path: Path) -> List[Dict[str, str]]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.S)
    blocks = re.findall(r'(?ms)^rule\s+"[^"]+".*?^end\s*$', text)
    rules: List[Dict[str, str]] = []
    for block in blocks:
        rule = parse_rule_block(block)
        if rule["code"]:
            rules.append(rule)
    return rules


def parse_impl_rule_blocks(path: Path) -> Dict[str, str]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.S)
    blocks = re.findall(r'(?ms)^rule\s+"[^"]+".*?^end\s*$', text)
    rules_by_code: Dict[str, str] = {}
    for block in blocks:
        code = extract_rule_code(block)
        if code:
            rules_by_code[code] = block.strip() + "\n"
    return rules_by_code


def sync_source_with_impl_rules(impl_path: Path = IMPL_FILE, source_dir: Path = DEFAULT_SOURCE_DIR) -> Dict[str, int]:
    source_dir.mkdir(parents=True, exist_ok=True)
    impl_rule_blocks = parse_impl_rule_blocks(impl_path)
    created = 0
    updated = 0
    unchanged = 0
    for code, block in impl_rule_blocks.items():
        # Normalizar el nombre de la regla a formato R{code}
        normalized_block = normalize_rule_block(block)
        source_file = source_dir / f"R{code}.drl"
        if not source_file.exists():
            source_file.write_text(normalized_block, encoding="utf-8")
            created += 1
            continue
        current = source_file.read_text(encoding="utf-8", errors="ignore")
        if current.strip() != normalized_block.strip():
            source_file.write_text(normalized_block, encoding="utf-8")
            updated += 1
        else:
            unchanged += 1
    return {
        "impl_rules_with_code": len(impl_rule_blocks),
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
    }


def parse_source_rules(source_dir: Path) -> Dict[str, Dict[str, str]]:
    source_rules: Dict[str, Dict[str, str]] = {}
    for file in sorted(source_dir.glob("R*.drl")):
        code_match = re.search(r"R(\d{4})\.drl$", file.name)
        if not code_match:
            continue
        code = code_match.group(1)
        text = file.read_text(encoding="utf-8", errors="ignore")
        block = text.strip()
        rule = parse_rule_block(block)
        rule["code"] = code
        rule["name"] = normalize_rule_name(extract_rule_name(block)) or f"R{code}"
        source_rules[code] = rule
    return source_rules


def similarity(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    a_tokens = set(a.split())
    b_tokens = set(b.split())
    if not a_tokens or not b_tokens:
        return 0.0
    inter = len(a_tokens & b_tokens)
    union = len(a_tokens | b_tokens)
    jaccard = inter / union if union else 0.0
    if a == b:
        return 1.0
    return jaccard


def associate_impl_with_source(impl_path: Path = IMPL_FILE, source_dir: Path = DEFAULT_SOURCE_DIR) -> Dict[str, Any]:
    impl_rules = parse_impl_rules(impl_path)
    source_rules = parse_source_rules(source_dir)
    matches: List[Dict[str, Any]] = []
    unmatched: List[Dict[str, Any]] = []
    direct_matches = 0
    description_matches = 0
    condition_matches = 0

    for rule in impl_rules:
        code = rule["code"]
        matched_source = source_rules.get(code)
        if matched_source:
            direct_matches += 1
            matches.append({
                "impl_code": code,
                "impl_rule": rule["name"],
                "impl_desc": rule["description"],
                "source_code": code,
                "source_rule": matched_source["name"],
                "match_type": "same_code",
                "confidence": 1.0,
                "status": "matched",
            })
            continue

        best_match: Optional[Tuple[str, float, Dict[str, str]]] = None
        best_reason = ""
        for source_code, source_rule_data in source_rules.items():
            desc_score = similarity(rule["description"], source_rule_data["description"])
            cond_score = similarity(rule["condition"], source_rule_data["condition"])
            score = max(desc_score, cond_score)
            if score > 0.4:
                reason = "description" if desc_score >= cond_score else "condition"
                if best_match is None or score > best_match[1]:
                    best_match = (source_code, score, source_rule_data)
                    best_reason = reason

        if best_match is not None:
            source_code, score, source_rule_data = best_match
            if best_reason == "description":
                description_matches += 1
                match_type = "description_similarity"
            else:
                condition_matches += 1
                match_type = "condition_similarity"
            matches.append({
                "impl_code": code,
                "impl_rule": rule["name"],
                "impl_desc": rule["description"],
                "source_code": source_code,
                "source_rule": source_rule_data["name"],
                "match_type": match_type,
                "confidence": round(score, 3),
                "status": "matched_by_similarity",
            })
        else:
            unmatched.append({
                "impl_code": code,
                "impl_rule": rule["name"],
                "impl_desc": rule["description"],
                "status": "unmatched_legacy_or_not_normalized",
            })

    return {
        "summary": {
            "impl_total_rules": len(impl_rules),
            "source_total_rules": len(source_rules),
            "direct_matches": direct_matches,
            "description_matches": description_matches,
            "condition_matches": condition_matches,
            "unmatched": len(unmatched),
        },
        "matches": matches,
        "unmatched": unmatched,
    }


def write_report(report: Dict[str, Any], report_dir: Path = DEFAULT_REPORTS_DIR) -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    output_path = report_dir / "rule_association_report.json"
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return output_path


def generate_excel_report(
    source_dir: Path = DEFAULT_SOURCE_DIR,
    generated_dir: Path = DEFAULT_GENERATED_DIR,
    output_path: Optional[Path] = None,
) -> Path:
    output_path = output_path or (DEFAULT_REPORTS_DIR / "reporte_reglas_completo.xlsx")
    doc_name_map = {
        "factura": "Factura",
        "boleta": "Boleta",
        "nota_de_credito": "Nota de crédito",
        "nota_de_debito": "Nota de débito",
        "retenciones": "Retenciones",
        "percepciones": "Percepciones",
        "resumen_diario": "Resumen Diario",
        "bajas": "Bajas",
    }
    report = associate_impl_with_source(IMPL_FILE, source_dir)
    impl_rules = parse_impl_rules(IMPL_FILE)
    impl_rule_by_code = {rule["code"]: rule for rule in impl_rules if rule.get("code")}
    source_codes = {path.stem[1:5] for path in source_dir.glob("R*.drl") if re.search(r"R\d{4}\.drl$", path.name)}
    impl_codes = set(impl_rule_by_code.keys())
    source_stubs = load_source_stubs()

    # Cargar mapeo de códigos a tipos de documento desde el Excel
    code_to_doc_types = load_code_to_doc_types()

    rule_to_docs: Dict[str, set] = defaultdict(set)
    for file in sorted(generated_dir.glob("*.drl")):
        doc_label = doc_name_map.get(file.stem, file.stem)
        text = file.read_text(encoding="utf-8", errors="ignore")
        # Captura tanto "R1001" como "1001_..." en el nombre de la regla
        for match in re.finditer(r'\brule\s+"(?:R)?(\d{4})(?:_[^"]*)?', text):
            rule_to_docs[match.group(1)].add(doc_label)

    workbook = Workbook()
    rule_sheet = workbook.active
    rule_sheet.title = "Reglas"
    headers = ["Nombre de la regla", "Tipo de documentos que la usan", "Descripcion de la regla", "Estado de implementacion"]
    rule_sheet.append(headers)

    rows: List[Tuple[str, str, str, str]] = []
    status_counter = Counter()
    status_by_doc: Dict[str, Counter] = defaultdict(Counter)

    for file in sorted(source_dir.glob("R*.drl")):
        code_match = re.search(r"R(\d{4})\.drl$", file.name)
        if not code_match:
            continue
        code = code_match.group(1)
        text = file.read_text(encoding="utf-8", errors="ignore")
        rule_name = normalize_rule_name(extract_rule_name(text)) or f"R{code}"
        desc_match = re.search(r'add(?:Error|Observacion|Excepcion)\("\d{4}"\s*,\s*"([^"]*)"\)', text)
        description = desc_match.group(1).strip() if desc_match else ""
        # Primero buscar en los archivos generados, luego en el Excel
        docs = sorted(rule_to_docs.get(code, set()))
        if not docs:
            docs = sorted(code_to_doc_types.get(code, set()))
        doc_text = ", ".join(docs) if docs else "No identificado"
        # Una regla esta implementada si existe en el DRL de implementacion, o si
        # su archivo source dejo de ser identico a la plantilla generada.
        # Sin plantilla de referencia no se puede afirmar nada del source: se
        # asume stub y decide unicamente el DRL.
        stub = source_stubs.get(code)
        es_stub = stub is None or text.strip() == stub.strip()
        status = "Implementada" if (code in impl_codes or not es_stub) else "Pendiente"
        rows.append((rule_name, doc_text, description, status))
        status_counter[status] += 1
        for doc in docs or ["No identificado"]:
            status_by_doc[doc][status] += 1

    for row in rows:
        rule_sheet.append(list(row))

    # Styling header
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in rule_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rule_sheet.column_dimensions["A"].width = 24
    rule_sheet.column_dimensions["B"].width = 58
    rule_sheet.column_dimensions["C"].width = 95
    rule_sheet.column_dimensions["D"].width = 18
    for row in rule_sheet.iter_rows(min_row=2, max_col=4):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        row[3].alignment = Alignment(horizontal="center", vertical="top")

    summary_sheet = workbook.create_sheet("Resumen")
    summary_sheet.append(["Metrica", "Valor"])
    summary_sheet.append(["Total reglas", len(rows)])
    summary_sheet.append(["Implementadas", status_counter["Implementada"]])
    summary_sheet.append(["Pendientes", status_counter["Pendiente"]])
    summary_sheet.append(["Codigos en el DRL de implementacion", len(impl_codes)])
    summary_sheet.append(["Codigos implementados sin regla source", len(impl_codes - source_codes)])
    summary_sheet.append([])
    per_doc_header_row = summary_sheet.max_row + 1
    summary_sheet.append(["Tipo de documento", "Implementadas", "Pendientes", "Total"])
    for doc in sorted(status_by_doc.keys()):
        imp = status_by_doc[doc]["Implementada"]
        pen = status_by_doc[doc]["Pendiente"]
        summary_sheet.append([doc, imp, pen, imp + pen])

    for cell in summary_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for cell in summary_sheet[per_doc_header_row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    summary_sheet.column_dimensions["A"].width = 35
    summary_sheet.column_dimensions["B"].width = 16
    summary_sheet.column_dimensions["C"].width = 14
    summary_sheet.column_dimensions["D"].width = 10

    for name, filtered in [("Implementadas", [r for r in rows if r[3] == "Implementada"]), ("Pendientes", [r for r in rows if r[3] == "Pendiente"])]:
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in filtered:
            ws.append(list(row))
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 58
        ws.column_dimensions["C"].width = 95
        ws.column_dimensions["D"].width = 18
        for row in ws.iter_rows(min_row=2, max_col=4):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[2].alignment = Alignment(wrap_text=True, vertical="top")
            row[3].alignment = Alignment(horizontal="center", vertical="top")

    dup_ws = workbook.create_sheet("Reglas_duplicadas")
    dup_ws.append(["Codigo_impl", "Regla_impl", "Codigo_source", "Regla_source", "Tipo_coincidencia", "Confianza", "Razon"])
    # Codigos emitidos por mas de una regla del DRL de implementacion.
    code_counts = Counter(rule["code"] for rule in impl_rules if rule.get("code"))
    for code, count in sorted(code_counts.items()):
        if count < 2:
            continue
        nombres = [rule["name"] for rule in impl_rules if rule.get("code") == code]
        dup_ws.append([
            code,
            ", ".join(nombres),
            code,
            f"R{code}" if code in source_codes else "",
            "same_code_multiple_rules",
            1.0,
            f"El codigo {code} es emitido por {count} reglas distintas del DRL de implementacion.",
        ])

    similar_matches = [m for m in report.get("matches", []) if m.get("match_type") in {"description_similarity", "condition_similarity"}]
    for match in similar_matches:
        dup_ws.append([
            match.get("impl_code", ""),
            match.get("impl_rule", ""),
            match.get("source_code", ""),
            match.get("source_rule", ""),
            match.get("match_type", ""),
            match.get("confidence", 0),
            "Se detecta una coincidencia funcional con otra regla del source y probablemente corresponde a una regla duplicada o normalizada.",
        ])

    for cell in dup_ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dup_ws.column_dimensions["A"].width = 12
    dup_ws.column_dimensions["B"].width = 40
    dup_ws.column_dimensions["C"].width = 12
    dup_ws.column_dimensions["D"].width = 40
    dup_ws.column_dimensions["E"].width = 26
    dup_ws.column_dimensions["F"].width = 12
    dup_ws.column_dimensions["G"].width = 88
    for row in dup_ws.iter_rows(min_row=2, max_col=7):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[6].alignment = Alignment(wrap_text=True, vertical="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    except PermissionError:
        fallback = DEFAULT_REPORTS_DIR / f"reporte_reglas_completo_{len(list(DEFAULT_REPORTS_DIR.glob('reporte_reglas_completo*.xlsx')))+1}.xlsx"
        workbook.save(fallback)
        return fallback
    return output_path


def main() -> None:
    ensure_dirs()
    source_summary = generate_source_rules(DEFAULT_EXCEL, DEFAULT_SOURCE_DIR)
    sync_summary = sync_source_with_impl_rules(IMPL_FILE, DEFAULT_SOURCE_DIR)
    generated_summary = generate_generated_files(DEFAULT_SOURCE_DIR, DEFAULT_GENERATED_DIR)
    report = associate_impl_with_source(IMPL_FILE, DEFAULT_SOURCE_DIR)
    report_path = write_report(report, DEFAULT_REPORTS_DIR)
    excel_path = generate_excel_report(DEFAULT_SOURCE_DIR, DEFAULT_GENERATED_DIR, DEFAULT_REPORTS_DIR / "reporte_reglas_completo.xlsx")
    print(json.dumps({
        "source_generation": source_summary,
        "source_sync_with_impl": sync_summary,
        "generated_files": generated_summary,
        "association": report["summary"],
        "report": str(report_path),
        "excel": str(excel_path),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
